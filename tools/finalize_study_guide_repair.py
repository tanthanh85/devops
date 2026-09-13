from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

PATH = "/Users/thandoan/Documents/Presentations/DevOps/Network-DevOps-Professional-Study-Guide.docx"


def shade(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = pPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd"); pPr.append(shd)
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), fill)


def borders(p, top=False, bottom=False):
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr"); pPr.append(pBdr)
    for side in (["top"] if top else []) + (["bottom"] if bottom else []) + ["left", "right"]:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8" if side == "left" else "4")
        el.set(qn("w:space"), "5")
        el.set(qn("w:color"), "2B7EA1" if side == "left" else "9ECAE1")
        pBdr.append(el)


doc = Document(PATH)
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == "**WHY THIS MATTERS TO A NETWORK ENGINEER**":
        p.text = "Why this matters to a network engineer"
        for r in p.runs:
            r.bold = True; r.font.name = "Arial"; r.font.size = Pt(10); r.font.color.rgb = RGBColor(0, 76, 135)
        p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
        shade(p, "EAF3F8"); borders(p, top=True)
        if i + 1 < len(doc.paragraphs):
            b = doc.paragraphs[i + 1]
            shade(b, "F5FAFC"); borders(b, bottom=True)
            b.paragraph_format.left_indent = Mm(4); b.paragraph_format.right_indent = Mm(3)
    if p.text.strip().startswith("##### 2.10.7.1 Practical control chain"):
        p.text = "Practical control chain for a protected deployment"
        p.style = doc.styles["Heading 4"]

doc.save(PATH)
