from collections import defaultdict, deque
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PATH = "/Users/thandoan/Documents/Presentations/DevOps/Network-DevOps-Professional-Study-Guide.docx"
MODULES = {
    "Network Automation Review",
    "Introducing the DevOps Model",
    "Infrastructure as Code and On-Demand Environments",
    "Packaging and Operating Applications",
    "Continuous Integration Delivery and Deployment Validation",
    "Security and Observability",
}


def add_level(abstract, level, start, text, indent, hanging):
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), str(level))
    start_el = OxmlElement("w:start"); start_el.set(qn("w:val"), str(start)); lvl.append(start_el)
    fmt = OxmlElement("w:numFmt"); fmt.set(qn("w:val"), "decimal"); lvl.append(fmt)
    if level > 0:
        restart = OxmlElement("w:lvlRestart"); restart.set(qn("w:val"), str(level)); lvl.append(restart)
    lvl_text = OxmlElement("w:lvlText"); lvl_text.set(qn("w:val"), text); lvl.append(lvl_text)
    suff = OxmlElement("w:suff"); suff.set(qn("w:val"), "space"); lvl.append(suff)
    jc = OxmlElement("w:lvlJc"); jc.set(qn("w:val"), "left"); lvl.append(jc)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab"); tab.set(qn("w:val"), "num"); tab.set(qn("w:pos"), str(indent)); tabs.append(tab)
    ind = OxmlElement("w:ind"); ind.set(qn("w:left"), str(indent)); ind.set(qn("w:hanging"), str(hanging))
    ppr.extend([tabs, ind]); lvl.append(ppr)
    abstract.append(lvl)


doc = Document(PATH)
numbering = doc.part.numbering_part.element
abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
abstract_id = max(abstract_ids or [0]) + 1
num_id = max(num_ids or [0]) + 1

abstract = OxmlElement("w:abstractNum")
abstract.set(qn("w:abstractNumId"), str(abstract_id))
multi = OxmlElement("w:multiLevelType"); multi.set(qn("w:val"), "multilevel"); abstract.append(multi)
add_level(abstract, 0, 0, "%1", 0, 0)
add_level(abstract, 1, 1, "%1.%2", 540, 540)
add_level(abstract, 2, 1, "%1.%2.%3", 720, 720)
numbering.insert(0, abstract)

num = OxmlElement("w:num"); num.set(qn("w:numId"), str(num_id))
aid = OxmlElement("w:abstractNumId"); aid.set(qn("w:val"), str(abstract_id)); num.append(aid)
numbering.append(num)

in_modules = False
toc_numbers = defaultdict(deque)
module_number = -1
major_number = 0
minor_number = 0
for p in doc.paragraphs:
    if p.style.name in {"Heading 1", "Heading 2", "Heading 3"}:
        ppr = p._p.get_or_add_pPr()
        old = ppr.find(qn("w:numPr"))
        if old is not None:
            ppr.remove(old)
    if p.text.strip() in MODULES and p.style.name == "Heading 1":
        in_modules = True
        module_number += 1
        major_number = 0
        minor_number = 0
    if (not in_modules or p.style.name not in {"Heading 1", "Heading 2", "Heading 3"}
            or p.text.strip() == "Learning objectives"):
        continue
    level = {"Heading 1": 0, "Heading 2": 1, "Heading 3": 2}[p.style.name]
    if level == 0:
        section_number = str(module_number)
    elif level == 1:
        major_number += 1
        minor_number = 0
        section_number = f"{module_number}.{major_number}"
    else:
        minor_number += 1
        section_number = f"{module_number}.{major_number}.{minor_number}"
    toc_numbers[p.text.strip()].append(section_number)
    ppr = p._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), str(level))
    nid = OxmlElement("w:numId"); nid.set(qn("w:val"), str(num_id))
    numpr.extend([ilvl, nid]); ppr.insert(0, numpr)

# Keep the cached TOC readable in renderers that do not automatically copy
# direct paragraph numbering into lower-level TOC entries. Hyperlinks and page
# reference fields are preserved; only the visible title text is prefixed.
for p in doc.paragraphs:
    if not p.style.name.lower().startswith("toc") or "\t" not in p.text:
        continue
    title = p.text.split("\t", 1)[0].strip()
    bare = re.sub(r"^\d+(?:\.\d+)*\s+", "", title)
    if toc_numbers[bare]:
        number = toc_numbers[bare].popleft()
        texts = p._p.xpath(".//w:t")
        if texts:
            texts[0].text = re.sub(r"^\d+(?:\.\d+)*\s+", "", texts[0].text or "")
            texts[0].text = f"{number} {texts[0].text}"

doc.save(PATH)
print(PATH)
