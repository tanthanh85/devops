#!/usr/bin/env python3
from pathlib import Path
import re, html
from docx import Document
from docx.shared import Mm, Pt, RGBColor, Inches
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
THEORY = ROOT / "theory"
OUT = ROOT / "Network-DevOps-Professional-Study-Guide.docx"
NAVY, BLUE, PALE, GRAY, DARK = "12345B", "1877B8", "EAF3F8", "D9E1E8", "263746"

MODULES = [
    (0, "Network Automation Review", "Can we automate and verify the network outcome?", "module-00-network-automation-review.md", "INTENT  →  LOGIC  →  DEVICE  →  EVIDENCE"),
    (1, "Introducing the DevOps Model", "Can a team own, deliver, measure, and improve the system?", "module-01-devops-model.md", "FLOW  →  FEEDBACK  →  EVIDENCE  →  IMPROVEMENT"),
    (2, "Infrastructure as Code and On-Demand Environments", "Can we recreate and govern the environment?", "module-02-infrastructure-as-code.md", "PLAN  →  CREATE  →  CONFIGURE  →  VERIFY  →  DESTROY"),
    (3, "Packaging and Operating Applications", "Can we reproduce and operate the application runtime?", "module-03-packaging-applications.md", "SOURCE  →  IMAGE  →  RUNTIME  →  PLATFORM"),
    (4, "Continuous Integration Delivery and Deployment Validation", "Can we automate the delivery process safely?", "module-04-cicd-delivery-validation.md", "SOURCE  →  TEST  →  ARTIFACT  →  PROMOTE  →  VERIFY"),
    (5, "Security and Observability", "Can we trust the system and explain what happened?", "module-05-security-observability.md", "IDENTITY  →  CONTROL  →  TELEMETRY  →  RESPONSE"),
]

FALLBACK_OBJECTIVES = {
    1: ["Explain the DevOps operating model and CALMS.", "Relate flow, feedback, evidence, and shared ownership to reliable delivery.", "Compare manual delivery, ad hoc automation, and governed pipelines."],
    2: ["Separate infrastructure provisioning, configuration, and application logic.", "Explain Terraform state, drift, lifecycle, and safe cleanup.", "Design an on-demand environment that hands explicit outputs to configuration automation."],
    3: ["Build and assess reproducible container images.", "Explain multitier application contracts, Docker Compose, and failure boundaries.", "Select Kubernetes only when its orchestration capabilities are justified."],
    4: ["Design purpose-specific CI/CD pipelines with defensible gates.", "Build once, identify artifacts, and promote without rebuilding.", "Validate deployment outcomes and choose an appropriate recovery response."],
    5: ["Define trust boundaries for source, runners, identities, artifacts, and targets.", "Correlate logs, metrics, traces, events, and deployment evidence.", "Use operational feedback to improve security, stability, and delivery."],
}

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn("w:shd"))
    if shd is None: shd = OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"), color)

def borders(table, color=GRAY):
    tblPr=table._tbl.tblPr; b=tblPr.find(qn("w:tblBorders"))
    if b is None: b=OxmlElement("w:tblBorders"); tblPr.append(b)
    for edge in ("top","left","bottom","right","insideH","insideV"):
        e=OxmlElement(f"w:{edge}"); e.set(qn("w:val"),"single"); e.set(qn("w:sz"),"5"); e.set(qn("w:color"),color); b.append(e)

def set_cell_margin(cell, top=100, start=120, bottom=100, end=120):
    tc=cell._tc; pr=tc.get_or_add_tcPr(); mar=pr.first_child_found_in("w:tcMar")
    if mar is None: mar=OxmlElement("w:tcMar"); pr.append(mar)
    for name,val in (("top",top),("start",start),("bottom",bottom),("end",end)):
        x=OxmlElement(f"w:{name}"); x.set(qn("w:w"),str(val)); x.set(qn("w:type"),"dxa"); mar.append(x)

def field(paragraph, instruction, display=""):
    r=paragraph.add_run(); begin=OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"),"begin")
    instr=OxmlElement("w:instrText"); instr.set(qn("xml:space"),"preserve"); instr.text=instruction
    sep=OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"),"separate")
    end=OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"),"end")
    r._r.append(begin)
    paragraph.add_run()._r.append(instr)
    paragraph.add_run()._r.append(sep)
    paragraph.add_run(display)
    paragraph.add_run()._r.append(end)

def page_break(doc): doc.add_page_break()

def configure(doc):
    sec=doc.sections[0]; sec.page_height=Mm(297); sec.page_width=Mm(210); sec.top_margin=Mm(21); sec.bottom_margin=Mm(21); sec.left_margin=Mm(24); sec.right_margin=Mm(19); sec.gutter=Mm(0)
    normal=doc.styles["Normal"]; normal.font.name="Aptos"; normal.font.size=Pt(10.3); normal.font.color.rgb=RGBColor.from_string(DARK)
    normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.13; normal.paragraph_format.widow_control=True
    for name,size,color,space in (("Title",32,"000000",14),("Subtitle",15,"000000",10),("Heading 1",23,"000000",16),("Heading 2",16,NAVY,10),("Heading 3",12,BLUE,7),("Heading 4",10.5,DARK,5)):
        st=doc.styles[name]; st.font.name="Aptos Display"; st.font.size=Pt(size); st.font.bold=True if name!="Subtitle" else False; st.font.color.rgb=RGBColor.from_string(color)
        st.paragraph_format.space_before=Pt(space); st.paragraph_format.space_after=Pt(5); st.paragraph_format.keep_with_next=True
        pPr=st.element.get_or_add_pPr(); old=pPr.find(qn("w:pBdr"))
        if old is not None: pPr.remove(old)
    code=doc.styles.add_style("Code Block",WD_STYLE_TYPE.PARAGRAPH); code.font.name="Consolas"; code.font.size=Pt(8.6); code.font.color.rgb=RGBColor.from_string("1F2933"); code.paragraph_format.left_indent=Mm(4); code.paragraph_format.right_indent=Mm(3); code.paragraph_format.space_before=Pt(4); code.paragraph_format.space_after=Pt(7); code.paragraph_format.keep_together=True
    cap=doc.styles["Caption"]; cap.font.name="Aptos"; cap.font.size=Pt(8.5); cap.font.italic=True; cap.font.color.rgb=RGBColor.from_string("465968"); cap.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER; cap.paragraph_format.keep_with_next=True
    for sec in doc.sections:
        sec.header_distance=Mm(9); sec.footer_distance=Mm(10)

def inline(p,text,bold=False):
    # lightweight Markdown inline formatting
    parts=re.split(r'(\*\*.*?\*\*|`.*?`|\[.*?\]\(.*?\))',text)
    for x in parts:
        if not x: continue
        if x.startswith("**") and x.endswith("**"): r=p.add_run(x[2:-2]); r.bold=True
        elif x.startswith("`") and x.endswith("`"): r=p.add_run(x[1:-1]); r.font.name="Consolas"; r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(NAVY)
        elif re.match(r'^\[.*\]\(.*\)$',x):
            m=re.match(r'^\[(.*)\]\((.*)\)$',x); p.add_run(m.group(1))
        else: p.add_run(html.unescape(re.sub(r'<[^>]+>','',x)))

def add_caption(doc,label,title):
    p=doc.add_paragraph(style="Caption"); p.paragraph_format.keep_with_next=False
    p.add_run(label+" "); field(p,f"SEQ {label} \\* ARABIC","1"); p.add_run(". "+title)

def add_table(doc, rows, caption=None):
    if caption: add_caption(doc,"Table",caption)
    t=doc.add_table(rows=0, cols=len(rows[0])); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=True; borders(t)
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        for j,val in enumerate(row):
            cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margin(cells[j]); p=cells[j].paragraphs[0]; inline(p,val.strip())
            for run in p.runs: run.font.size=Pt(8.8)
            if i==0:
                shade(cells[j],NAVY)
                for run in p.runs: run.bold=True; run.font.color.rgb=RGBColor(255,255,255)
            elif i%2==0: shade(cells[j],"F3F7FA")
    # repeating header
    trPr=t.rows[0]._tr.get_or_add_trPr(); hdr=OxmlElement("w:tblHeader"); hdr.set(qn("w:val"),"true"); trPr.append(hdr)
    return t

def add_figure(doc,path,alt,module):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next=True
    try:
        from PIL import Image
        with Image.open(path) as im: w,h=im.size
        maxw,maxh=6.45,4.75; width=min(maxw,maxh*w/h)
        p.add_run().add_picture(str(path),width=Inches(width))
        add_caption(doc,"Figure",alt.rstrip("."))
    except Exception:
        q=doc.add_paragraph(f"Figure unavailable: {alt}"); q.style="Caption"

def extract_objectives(text):
    m=re.search(r'## 2\. Learning objectives\s+(.*?)(?=\n## )',text,re.S|re.I)
    if not m: return []
    return [re.sub(r'^-\s*','',x).strip() for x in m.group(1).splitlines() if x.startswith('- ')][:6]

def module_opener(doc,num,title,question,focus,text):
    if num>0: doc.add_section(WD_SECTION.ODD_PAGE)
    p=doc.add_paragraph(); p.add_run(f"MODULE {num}").bold=True; p.runs[0].font.color.rgb=RGBColor.from_string(BLUE); p.runs[0].font.size=Pt(13)
    h=doc.add_paragraph(title,style="Title"); h.paragraph_format.space_before=Pt(22)
    p=doc.add_paragraph(); r=p.add_run("Central engineering question\n"); r.bold=True; r.font.color.rgb=RGBColor.from_string(BLUE); p.add_run(question)
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; borders(t,BLUE); shade(t.cell(0,0),PALE); set_cell_margin(t.cell(0,0),180,180,180,180)
    p=t.cell(0,0).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(focus); r.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY)
    doc.add_paragraph("Learning objectives",style="Heading 2")
    objectives=extract_objectives(text) or FALLBACK_OBJECTIVES[num]
    for x in objectives: inline(doc.add_paragraph(style="List Bullet"),x)
    meta=doc.add_table(rows=4,cols=2); borders(meta); meta.alignment=WD_TABLE_ALIGNMENT.CENTER
    values=(("Estimated theory time",{0:"2 hours",1:"3 hours",2:"3 hours",3:"4.5 hours",4:"4.5 hours",5:"3 hours"}[num]),("Lab connection","Apply the module controls to the cumulative application and delivery system"),("Core classroom topics","Concepts, decisions, workflows, and required controls"),("Advanced reference","Detailed implementation and platform considerations"))
    for i,(a,b) in enumerate(values): meta.cell(i,0).text=a; meta.cell(i,1).text=b; shade(meta.cell(i,0),PALE); meta.cell(i,0).paragraphs[0].runs[0].bold=True
    page_break(doc)

def parse_markdown(doc,path,module):
    lines=path.read_text().splitlines(); i=0; in_code=False; code=[]; skip=False
    while i<len(lines):
        line=lines[i]
        if line.startswith("# "): i+=1; continue
        if line.startswith("## 1. Purpose") or line.startswith("## 2. Learning objectives"): skip=True; i+=1; continue
        if skip and line.startswith("## "): skip=False
        if skip: i+=1; continue
        if line.startswith("```"):
            if not in_code: in_code=True; code=[]
            else:
                p=doc.add_paragraph(style="Code Block"); p.paragraph_format.keep_together=len(code)<24; r=p.add_run("\n".join(code)); r.font.name="Consolas"
                shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),"F2F4F6"); p._p.get_or_add_pPr().append(shd); in_code=False
            i+=1; continue
        if in_code: code.append(line); i+=1; continue
        if line.startswith("<p align="):
            blob=" ".join(lines[i:i+3]); m=re.search(r'<img src="([^"]+)" alt="([^"]+)"',blob)
            if m: add_figure(doc,(path.parent/m.group(1)).resolve(),m.group(2),module)
            while i<len(lines) and "</p>" not in lines[i]: i+=1
            i+=1; continue
        m=re.match(r'^(#{2,4})\s+(.*)$',line)
        if m:
            level=len(m.group(1))-1; title=re.sub(r'^\d+(?:\.\d+)*\.?\s*','',m.group(2)).strip()
            doc.add_paragraph(title,style=f"Heading {level}"); i+=1; continue
        if line.startswith("|") and i+1<len(lines) and re.match(r'^\|[\s:|-]+\|$',lines[i+1]):
            rows=[]; i+=2
            header=[x.strip() for x in line.strip('|').split('|')]
            rows.append(header)
            while i<len(lines) and lines[i].startswith('|'):
                rows.append([x.strip() for x in lines[i].strip('|').split('|')]); i+=1
            add_table(doc,rows,caption="Technical comparison and reference data"); continue
        if re.match(r'^[-*] ',line):
            p=doc.add_paragraph(style="List Bullet"); inline(p,line[2:]); i+=1; continue
        if re.match(r'^\d+\. ',line):
            p=doc.add_paragraph(style="List Number"); inline(p,re.sub(r'^\d+\.\s*','',line)); i+=1; continue
        if line.startswith("> "):
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Mm(7); r=p.add_run(line[2:]); r.italic=True; r.font.color.rgb=RGBColor.from_string(NAVY); i+=1; continue
        if not line.strip() or line.strip()=="---": i+=1; continue
        # join normal paragraph lines
        para=[line.strip()]; i+=1
        while i<len(lines) and lines[i].strip() and not re.match(r'^(#{1,4})\s|^```|^[-*] |^\d+\. |^\||^<p align=|^> ',lines[i]): para.append(lines[i].strip()); i+=1
        inline(doc.add_paragraph()," ".join(para))

def front_matter(doc):
    # Cover
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(45); r=p.add_run("IMPLEMENTING DEVOPS SOLUTIONS\nAND PRACTICES"); r.bold=True; r.font.name="Aptos Display"; r.font.size=Pt(28); r.font.color.rgb=RGBColor.from_string(NAVY)
    p=doc.add_paragraph("using Cisco Platforms (DevOps) v1.0",style="Title"); p.runs[0].font.size=Pt(21)
    p=doc.add_paragraph("Professional Study Guide",style="Subtitle"); p.runs[0].font.size=Pt(16); p.runs[0].bold=True
    doc.add_paragraph("Engineering Reliable Delivery Systems for Network Automation",style="Subtitle")
    doc.add_paragraph("\nINTENT  →  AUTOMATION  →  DELIVERY PIPELINE\n↓\nCONTAINER PLATFORM  →  APPLICATION AND NETWORK\n↓\nOBSERVABILITY  →  EVIDENCE",style=None).alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(70); p.add_run("Independent Study Guide\nNot an official Cisco publication").bold=True
    doc.add_paragraph("Author / Instructor: ______________________________\nVersion 1.0  |  2026")
    page_break(doc)
    doc.add_paragraph("Implementing DevOps Solutions and Practices using Cisco Platforms (DevOps) v1.0",style="Title"); doc.add_paragraph("Professional Study Guide\nEngineering Reliable Delivery Systems for Network Automation",style="Subtitle"); doc.add_paragraph("Five-Day Instructor-Led Course\n\nVersion 1.0\n2026\n\nAuthor / Instructor: ______________________________\n\nIndependent publication. Not an official Cisco publication."); page_break(doc)
    doc.add_paragraph("Publication notice",style="Heading 1"); doc.add_paragraph("This guide is independently authored and is not an official Cisco publication. Cisco and other product names may be trademarks of their respective owners. Product behavior varies by version and platform. Validate all commands and examples in an authorized environment. Network changes can affect production services; use approved test systems, change controls, credentials, and recovery procedures."); page_break(doc)
    doc.add_paragraph("About this guide",style="Heading 1"); doc.add_paragraph("This guide develops a controlled software-delivery system around network automation. It begins with familiar automation building blocks, then adds team ownership, Infrastructure as Code, reproducible application packaging, container platforms, CI/CD, security, observability, and operational feedback.")
    doc.add_paragraph("Who should read this guide",style="Heading 2"); doc.add_paragraph("Network and infrastructure professionals who can automate technical work and now need to engineer, deliver, and operate that automation reliably as a team.")
    doc.add_paragraph("Prerequisite knowledge",style="Heading 2"); doc.add_paragraph("Knowledge equivalent to DEVASC and DEVCOR is expected, including Python, Ansible, APIs, JSON and YAML, Git, and Linux command-line operation.")
    doc.add_paragraph("How to use this guide",style="Heading 2"); doc.add_paragraph("Read modules in sequence. Use diagrams to establish the system model, examples to examine implementation choices, and knowledge checks to test engineering judgment. Advanced material may be revisited as reference.")
    doc.add_paragraph("Course engineering journey",style="Heading 1"); add_figure(doc,THEORY/"assets/course-figures/course-engineering-journey.png","Course engineering journey from network automation to continuous improvement",0)
    doc.add_paragraph("Five-day course map",style="Heading 1")
    add_table(doc,[["Day","Engineering focus","Outcome"],["1","Automation review and DevOps model","Shared delivery and evidence model"],["2","Infrastructure as Code and packaging","Reproducible environment and image"],["3","Multitier operation and Kubernetes","Controlled application runtime"],["4","CI/CD and deployment validation","Qualified and promoted artifacts"],["5","Security and observability","Trusted operation and feedback"]],"Five-day learning sequence")
    doc.add_paragraph("Course reference architecture",style="Heading 1"); add_figure(doc,THEORY/"assets/diagrams/course-reference-architecture.png","Course reference delivery architecture and trust boundaries",0)
    doc.add_paragraph("Course reference scenario",style="Heading 1"); doc.add_paragraph("A consistent change scenario provides concrete names and values for engineering discussions. It is an example, not a production authorization.")
    add_table(doc,[["Context","Value"],["Site","campus-west"],["Target","distribution-01"],["Peer","routing-peer-01"],["Service","VLAN 120 USERS"],["Gateway","10.20.120.1/24"],["Routing","OSPF process 100, area 0"],["Change","CHG-2026-0042"]],"Course reference scenario")
    doc.add_paragraph("Module overview",style="Heading 1"); add_table(doc,[["Module","Subject","Central question"]]+[[str(n),t,q] for n,t,q,_,_ in MODULES],"Module engineering questions")
    for title,instr in (("Table of contents",'TOC \\o "1-3" \\h \\z \\u'),("List of figures",'TOC \\c "Figure"'),("List of tables",'TOC \\c "Table"')):
        page_break(doc); doc.add_paragraph(title,style="Heading 1"); p=doc.add_paragraph(); field(p,instr,"Update this field in Microsoft Word")

def headers(doc):
    # Sections inherit the same linked header/footer, so author each story once.
    sec=doc.sections[0]
    hp=sec.header.paragraphs[0]; hp.text="IMPLEMENTING DEVOPS SOLUTIONS AND PRACTICES  |  STUDY GUIDE"; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    for r in hp.runs: r.font.name="Aptos"; r.font.size=Pt(7.5); r.font.color.rgb=RGBColor.from_string("667788")
    fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; field(fp,"PAGE","1")

def main():
    doc=Document(); configure(doc); front_matter(doc)
    for num,title,q,file,focus in MODULES:
        text=(THEORY/file).read_text(); module_opener(doc,num,title,q,focus,text); parse_markdown(doc,THEORY/file,num)
    headers(doc)
    # Keep TOC and caption fields refreshable, but do not force an update when
    # opening the file. Word otherwise presents an external-field security
    # warning even though every figure is embedded in the document package.
    doc.core_properties.title="Implementing DevOps Solutions and Practices using Cisco Platforms (DevOps) v1.0"; doc.core_properties.subject="Professional study guide for engineering reliable delivery systems"; doc.core_properties.author="Independent Study Guide"
    doc.save(OUT); print(OUT)

if __name__=="__main__": main()
